#!/usr/bin/env python3
"""
make_config_xlsx.py  —  V2
Regenerates config/site-config.xlsx from the defaults in config_schema.py.

You normally do NOT need to run this. It exists so a blank, correctly
structured workbook can always be recreated if the original is lost.

    python3 tools/make_config_xlsx.py

Requires:  pip install openpyxl
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config_schema as S

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "config", "site-config.xlsx")

HEAD_FILL = PatternFill("solid", fgColor="264653")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11, name="Calibri")
KEY_FONT = Font(bold=True, size=10, name="Consolas")
VAL_FILL = PatternFill("solid", fgColor="FFF6EE")
HELP_FONT = Font(size=9, color="6E6E6E", italic=True)
THIN = Side(style="thin", color="DADADA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center")
        cell.border = BORDER
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def kv_sheet(wb, title, rows, note):
    ws = wb.create_sheet(title)
    ws.append(["Setting", "Value", "What it does"])
    style_header(ws, 3)
    for key, val, help_text in rows:
        ws.append([key, val, help_text])
        r = ws.max_row
        ws.cell(row=r, column=1).font = KEY_FONT
        ws.cell(row=r, column=2).fill = VAL_FILL
        ws.cell(row=r, column=3).font = HELP_FONT
        for c in range(1, 4):
            ws.cell(row=r, column=c).border = BORDER
            ws.cell(row=r, column=c).alignment = Alignment(vertical="center", wrap_text=(c == 3))
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 58
    ws.column_dimensions["C"].width = 78
    ws.append([])
    ws.append(["NOTE", note])
    ws.cell(row=ws.max_row, column=2).font = HELP_FONT
    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)

    # ---------- Read Me ----------
    ws = wb.create_sheet("Read Me")
    ws.append(["Pillai & Sons Motor Company — website configuration"])
    ws["A1"].font = Font(bold=True, size=16, color="264653")
    lines = [
        "",
        "Everything on this workbook drives the live website. Edit the Value column only.",
        "",
        "HOW TO APPLY YOUR CHANGES  —  pick either route, both do the same thing:",
        "",
        "  ROUTE A — no software needed (recommended)",
        "     1. Save this file.",
        "     2. Open  admin/index.html  in Chrome, Edge or Safari.",
        "     3. Go to the 'Excel Config' tab and drop this .xlsx file on it.",
        "     4. Press 'Apply & Save'. The site updates immediately.",
        "     5. Press 'Download site-config.js' and put that file in  assets/js/  ",
        "        before pushing to GitHub, so visitors get the new settings too.",
        "",
        "  ROUTE B — with Python installed",
        "     1. Save this file.",
        "     2. Run:   python3 tools/build_config.py",
        "     3. Commit the regenerated  assets/js/site-config.js  and push.",
        "",
        "SHEETS",
        "  Business  — name, phone, WhatsApp, email, address, Google Map, counters.",
        "  Colors    — every colour on the site. Use 6-digit hex like #24A19C.",
        "  Fonts     — font family, every font size, weight, spacing and radius.",
        "  Social    — social profile links. Set Enabled to 'no' to hide an icon.",
        "  Themes    — the palettes offered to visitors by the on-page colour switcher.",
        "",
        "TIPS",
        "  * Never rename a Setting key, or change the sheet names. Values only.",
        "  * Colours must be hex (#RRGGBB). Sizes are plain numbers, no 'px'.",
        "  * whatsapp_number must be country code + number, digits only: 919876543210.",
        "  * map_embed_url must be an EMBED url. In Google Maps: Share > Embed a map >",
        "    copy only the src=\"...\" part.",
    ]
    for line in lines:
        ws.append([line])
    ws.column_dimensions["A"].width = 100
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=1).alignment = Alignment(vertical="center")

    kv_sheet(wb, "Business", S.BUSINESS,
             "Phone numbers: digits only in 'phone' and 'whatsapp_number'. Pretty formatting goes in 'phone_display'.")
    kv_sheet(wb, "Colors", S.COLORS,
             "All values must be 6-digit hex codes starting with #.")
    kv_sheet(wb, "Fonts", S.FONTS,
             "Sizes are numbers only (no 'px'). Bundled offline families: Montserrat, Poppins.")

    # ---------- Social ----------
    ws = wb.create_sheet("Social")
    ws.append(["Platform", "URL", "Enabled", "Label"])
    style_header(ws, 4)
    for platform, url, enabled, label in S.SOCIAL:
        ws.append([platform, url, enabled, label])
        r = ws.max_row
        ws.cell(row=r, column=1).font = KEY_FONT
        ws.cell(row=r, column=2).fill = VAL_FILL
        ws.cell(row=r, column=3).fill = VAL_FILL
        for c in range(1, 5):
            ws.cell(row=r, column=c).border = BORDER
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 62
    ws.column_dimensions["C"].width = 11
    ws.column_dimensions["D"].width = 20
    ws.append([])
    ws.append(["NOTE", "Enabled = yes / no. Leave the WhatsApp URL blank to build it "
                       "automatically from whatsapp_number on the Business sheet."])
    ws.cell(row=ws.max_row, column=2).font = HELP_FONT

    # ---------- Themes ----------
    ws = wb.create_sheet("Themes")
    ws.append(["Theme Name", "Primary", "Primary Ink", "Secondary", "Tertiary",
               "Tertiary Ink", "Accent Cream", "Footer BG", "Top Bar BG"])
    style_header(ws, 9)
    for row in S.THEMES:
        ws.append(list(row))
        r = ws.max_row
        ws.cell(row=r, column=1).font = KEY_FONT
        for c in range(2, 10):
            ws.cell(row=r, column=c).fill = VAL_FILL
        for c in range(1, 10):
            ws.cell(row=r, column=c).border = BORDER
    ws.column_dimensions["A"].width = 22
    for c in range(2, 10):
        ws.column_dimensions[get_column_letter(c)].width = 16
    ws.append([])
    ws.append(["NOTE", "The first row is the site default. These palettes appear in the "
                       "colour switcher visitors can open from the floating paint icon. "
                       "The two 'Ink' columns are the darker shades used wherever small "
                       "text meets the colour — they are what keep the site readable."])
    ws.cell(row=ws.max_row, column=2).font = HELP_FONT

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
